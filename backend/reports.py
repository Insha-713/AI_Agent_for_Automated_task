"""
AgentFlow - CSV and Excel Report Module (Stage 4)
Automated data cleaning, KPI computation, Recharts-ready visualization generation,
executive summary authoring, and multi-format report export (Excel, CSV, Markdown).
Designed for Indian SMBs, job seekers, and students with zero paid APIs.
"""

import os
import re
import csv
import json
import math
from datetime import datetime, timezone, timedelta, date
from decimal import Decimal
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple, Union

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Workspace & Reports paths
PROJECT_ROOT = Path(__file__).resolve().parent.parent
WORKSPACE_PATH = (PROJECT_ROOT / os.getenv("WORKSPACE_DIR", "workspace")).resolve()
REPORTS_DIR = (WORKSPACE_PATH / "reports").resolve()
SAMPLES_DIR = (WORKSPACE_PATH / "samples").resolve()

REPORTS_DIR.mkdir(parents=True, exist_ok=True)
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

# IST Timezone helper
IST = timezone(timedelta(hours=5, minutes=30))


def clean_nan_and_inf(val: Any) -> Any:
    """
    Recursively cleans any dict, list, pandas, numpy, or primitive value:
    - NaN, +inf, -inf become None (null in JSON)
    - numpy ints and floats become normal Python numbers (or None if NaN/inf)
    - numpy bool_ becomes normal Python bool
    - pandas Timestamp and NaT become ISO strings or None
    - datetime / date become ISO strings
    - Decimal becomes float (or None if NaN/inf)
    - dict, list, tuple, set, numpy array, pandas Series/DataFrame are traversed and cleaned recursively
    """
    if val is None:
        return None

    # Handle pandas NaT
    if val is pd.NaT:
        return None

    # Handle floats & numpy floats
    if isinstance(val, (float, np.floating)):
        if math.isnan(val) or math.isinf(val) or np.isnan(val) or np.isinf(val):
            return None
        return float(val)

    # Handle numpy integer
    if isinstance(val, (np.integer,)):
        return int(val)

    # Handle numpy boolean
    if isinstance(val, (np.bool_,)):
        return bool(val)

    # Handle Decimal
    if isinstance(val, Decimal):
        if val.is_nan() or val.is_infinite():
            return None
        return float(val)

    # Handle pandas Timestamp
    if isinstance(val, pd.Timestamp):
        if pd.isna(val):
            return None
        return val.isoformat()

    # Handle datetime / date
    if isinstance(val, (datetime, date)):
        return val.isoformat()

    # Handle dictionaries
    if isinstance(val, dict):
        return {str(k): clean_nan_and_inf(v) for k, v in val.items()}

    # Handle lists, tuples, sets, numpy arrays
    if isinstance(val, (list, tuple, set, np.ndarray)):
        return [clean_nan_and_inf(item) for item in val]

    # Handle pandas Series / DataFrame
    if isinstance(val, pd.Series):
        return clean_nan_and_inf(val.to_dict())
    if isinstance(val, pd.DataFrame):
        return clean_nan_and_inf(val.to_dict(orient="records"))

    # Fallback check for pd.isna
    try:
        if pd.isna(val):
            return None
    except Exception:
        pass

    return val


def _clean_currency_or_numeric(val: Any) -> Any:
    """Helper to convert formatted currency strings (₹, $, commas, %) into float/int."""
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (float, np.floating)):
        if math.isnan(val) or math.isinf(val) or np.isnan(val) or np.isinf(val):
            return None
        return float(val)
    if isinstance(val, (int, np.integer)):
        return int(val)
    s = str(val).strip()
    if not s or s.lower() in ["none", "nan", "null", "nat", "n/a", "na"]:
        return None
    # Strip currency symbols and whitespace
    cleaned = re.sub(r"[₹$,\s]", "", s)
    if cleaned.endswith("%"):
        try:
            return float(cleaned[:-1])
        except ValueError:
            pass
    try:
        if "." in cleaned:
            f = float(cleaned)
            return None if (math.isnan(f) or math.isinf(f)) else f
        return int(cleaned)
    except ValueError:
        return s


def load_and_clean_tabular_data(file_path: Union[str, Path]) -> Dict[str, Any]:
    """
    Loads CSV or Excel data, strips whitespace, standardizes types,
    and returns cleaned records and metadata.
    """
    p = Path(file_path)
    if not p.is_absolute():
        p = (WORKSPACE_PATH / file_path).resolve()

    if not p.exists():
        # Check inside samples dir as fallback
        sample_candidate = SAMPLES_DIR / p.name
        if sample_candidate.exists():
            p = sample_candidate
        else:
            raise FileNotFoundError(f"File not found: {file_path}")

    # Read data using pandas
    suffix = p.suffix.lower()
    if suffix in [".xlsx", ".xls"]:
        df = pd.read_excel(p)
    elif suffix in [".csv", ".tsv", ".txt"]:
        try:
            df = pd.read_csv(p, encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(p, encoding="latin1")
    else:
        raise ValueError(f"Unsupported file format '{suffix}'. Supported: .csv, .xlsx, .xls")

    original_rows = len(df)
    original_cols = len(df.columns)

    # Standardize column headers: trim whitespace
    df.columns = [str(c).strip() for c in df.columns]

    # Clean numeric columns with formatting
    missing_counts = {}
    col_types = {}

    for col in df.columns:
        # Check missing count
        missing_count = int(df[col].isna().sum())
        missing_counts[col] = missing_count

        # Try converting string columns that look numeric
        if df[col].dtype == object:
            # Check if majority of non-null items can be numeric
            non_null = df[col].dropna()
            if len(non_null) > 0:
                converted = non_null.apply(_clean_currency_or_numeric)
                num_count = sum(isinstance(x, (int, float)) for x in converted)
                if num_count / len(non_null) > 0.7:
                    df[col] = df[col].apply(_clean_currency_or_numeric)

        # Detect column types
        if pd.api.types.is_numeric_dtype(df[col]):
            col_types[col] = "numeric"
        elif pd.api.types.is_datetime64_any_dtype(df[col]):
            col_types[col] = "datetime"
            df[col] = df[col].dt.strftime("%Y-%m-%d")
        else:
            # Check if strings are date-like (YYYY-MM-DD)
            valid_non_null = [str(v) for v in df[col].dropna() if str(v).strip() not in ["", "None", "nan", "null", "NaT"]]
            if valid_non_null and all(re.match(r"^\d{4}-\d{2}-\d{2}", str(v)) for v in valid_non_null[:5]):
                col_types[col] = "datetime"
            else:
                col_types[col] = "categorical"

    # Clean records properly without leaving NaN floats in dictionaries
    raw_records = df.to_dict(orient="records")
    cleaned_records = []
    for r in raw_records:
        cleaned_row = {}
        for c in df.columns:
            val = r.get(c)
            cleaned_row[c] = clean_nan_and_inf(val)
        cleaned_records.append(cleaned_row)

    return clean_nan_and_inf({
        "file_name": p.name,
        "file_path": str(p),
        "row_count": original_rows,
        "column_count": original_cols,
        "columns": list(df.columns),
        "column_types": col_types,
        "missing_counts": missing_counts,
        "records": cleaned_records,
    })


def calculate_kpis(cleaned_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Computes comprehensive KPIs across numeric, categorical, and temporal fields.
    Guarantees no NaN, Inf, or zero-division errors.
    """
    records = cleaned_data.get("records", [])
    col_types = cleaned_data.get("column_types", {})
    total_rows = cleaned_data.get("row_count", 0)
    col_count = cleaned_data.get("column_count", 0)

    if total_rows == 0:
        return {
            "dataset_summary": {
                "total_rows": 0,
                "total_columns": col_count,
                "data_quality_score": 0.0,
            },
            "numeric_metrics": {},
            "categorical_metrics": {},
            "temporal_metrics": {},
            "top_business_metrics": [],
            "status": "empty_dataset"
        }

    # Data Quality calculation (safe from zero-division)
    total_cells = total_rows * col_count
    total_missing = sum(cleaned_data.get("missing_counts", {}).values())
    if total_cells > 0:
        quality_score = max(0.0, round(((total_cells - total_missing) / total_cells) * 100, 1))
    else:
        quality_score = 0.0

    kpis: Dict[str, Any] = {
        "dataset_summary": {
            "total_rows": total_rows,
            "total_columns": col_count,
            "data_quality_score": quality_score,
        },
        "numeric_metrics": {},
        "categorical_metrics": {},
        "temporal_metrics": {},
        "top_business_metrics": [],
    }

    # Compute numeric KPIs with empty-column & zero-denominator guards
    for col, ctype in col_types.items():
        if ctype == "numeric":
            vals = [
                r[col] for r in records
                if r.get(col) is not None
                and isinstance(r.get(col), (int, float))
                and not math.isnan(r.get(col))
                and not math.isinf(r.get(col))
            ]
            if vals:
                total_val = sum(vals)
                avg_val = round(total_val / len(vals), 2)
                sorted_vals = sorted(vals)
                mid = len(sorted_vals) // 2
                median_val = (sorted_vals[mid] if len(sorted_vals) % 2 != 0 else (sorted_vals[mid - 1] + sorted_vals[mid]) / 2)
                min_val = min(vals)
                max_val = max(vals)

                kpis["numeric_metrics"][col] = {
                    "total": round(total_val, 2),
                    "average": avg_val,
                    "median": round(median_val, 2),
                    "min": min_val,
                    "max": max_val,
                    "count": len(vals),
                }
            else:
                kpis["numeric_metrics"][col] = {
                    "total": None,
                    "average": None,
                    "median": None,
                    "min": None,
                    "max": None,
                    "count": 0,
                }

    # Compute categorical KPIs with guards
    for col, ctype in col_types.items():
        if ctype == "categorical":
            counts: Dict[str, int] = {}
            for r in records:
                v = r.get(col)
                if v is not None:
                    s_val = str(v).strip()
                    if s_val and s_val.lower() not in ["none", "nan", "null", "nat", "n/a", "na"]:
                        counts[s_val] = counts.get(s_val, 0) + 1

            if counts:
                sorted_counts = sorted(counts.items(), key=lambda x: x[1], reverse=True)
                top_name, top_freq = sorted_counts[0]
                top_pct = round((top_freq / total_rows) * 100, 1) if total_rows > 0 else 0.0
                kpis["categorical_metrics"][col] = {
                    "unique_values_count": len(counts),
                    "top_value": top_name,
                    "top_value_frequency": top_freq,
                    "top_percentage": top_pct,
                    "distribution": dict(sorted_counts[:7]),
                }
            else:
                kpis["categorical_metrics"][col] = {
                    "unique_values_count": 0,
                    "top_value": "N/A",
                    "top_value_frequency": 0,
                    "top_percentage": 0.0,
                    "distribution": {},
                }

    # Compute temporal KPIs with guards
    for col, ctype in col_types.items():
        if ctype == "datetime":
            dates = [
                str(r[col]) for r in records
                if r.get(col) and str(r.get(col)).strip() not in ["None", "nan", "null", "NaT", ""]
            ]
            if dates:
                kpis["temporal_metrics"][col] = {
                    "earliest_date": min(dates),
                    "latest_date": max(dates),
                    "date_count": len(dates),
                }
            else:
                kpis["temporal_metrics"][col] = {
                    "earliest_date": "N/A",
                    "latest_date": "N/A",
                    "date_count": 0,
                }

    # Identify Top Business Metrics (Revenue, Units, Cost, CTC, etc.)
    num_metrics = kpis["numeric_metrics"]
    for col, stats in num_metrics.items():
        if stats["count"] == 0 or stats["total"] is None:
            continue
        lower_col = col.lower()
        if any(kw in lower_col for kw in ["revenue", "sales", "amount", "price", "income", "package", "cost", "expense", "profit", "units", "ctc"]):
            prefix = "₹" if any(kw in lower_col for kw in ["revenue", "sales", "amount", "price", "income", "cost", "expense", "profit", "inr"]) else ""
            suffix = " LPA" if ("ctc" in lower_col or "lpa" in lower_col) else ""
            formatted_total = f"{prefix}{stats['total']:,.2f}{suffix}" if prefix else f"{stats['total']:,}{suffix}"
            formatted_avg = f"{prefix}{stats['average']:,.2f}{suffix}" if prefix else f"{stats['average']:,}{suffix}"
            kpis["top_business_metrics"].append({
                "metric_name": col.replace("_", " ").title(),
                "total_display": formatted_total,
                "avg_display": formatted_avg,
                "raw_total": stats["total"],
                "raw_avg": stats["average"],
                "min": stats["min"],
                "max": stats["max"],
            })

    return clean_nan_and_inf(kpis)


def generate_charts_data(cleaned_data: Dict[str, Any], kpis: Dict[str, Any]) -> Dict[str, Any]:
    """
    Builds data structures tailored for Recharts (Bar, Line, Pie, Stats).
    Guarantees no NaN or Inf values.
    """
    records = cleaned_data.get("records", [])
    col_types = cleaned_data.get("column_types", {})
    numeric_cols = [c for c, t in col_types.items() if t == "numeric"]
    cat_cols = [c for c, t in col_types.items() if t == "categorical"]
    date_cols = [c for c, t in col_types.items() if t == "datetime"]

    # Choose primary numeric metric
    primary_num = None
    for kw in ["revenue", "sales", "package", "ctc", "amount", "expense", "cost", "score", "price", "units"]:
        for col in numeric_cols:
            if kw in col.lower() and kpis.get("numeric_metrics", {}).get(col, {}).get("count", 0) > 0:
                primary_num = col
                break
        if primary_num:
            break
    if not primary_num and numeric_cols:
        for col in numeric_cols:
            if kpis.get("numeric_metrics", {}).get(col, {}).get("count", 0) > 0:
                primary_num = col
                break

    # Choose primary categorical group
    primary_cat = None
    for kw in ["category", "branch", "department", "state", "city", "status", "role", "type"]:
        for col in cat_cols:
            if kw in col.lower():
                primary_cat = col
                break
        if primary_cat:
            break
    if not primary_cat and cat_cols:
        primary_cat = cat_cols[0]

    # 1. Bar Chart: Primary Category vs Sum of Primary Numeric
    bar_chart_data = []
    if primary_cat and primary_num:
        cat_agg: Dict[str, float] = {}
        for r in records:
            cat_raw = r.get(primary_cat)
            if cat_raw is None or str(cat_raw).strip() in ["", "None", "nan", "null"]:
                cat_val = "Unspecified"
            else:
                cat_val = str(cat_raw).strip()

            num_val = r.get(primary_num)
            if num_val is not None and isinstance(num_val, (int, float)):
                if not (math.isnan(num_val) or math.isinf(num_val)):
                    cat_agg[cat_val] = cat_agg.get(cat_val, 0.0) + float(num_val)

        sorted_cat = sorted(cat_agg.items(), key=lambda x: x[1], reverse=True)[:8]
        bar_chart_data = [
            {"name": k, "value": round(v, 2), "metric": primary_num.replace("_", " ").title()}
            for k, v in sorted_cat
            if v is not None and not math.isnan(v)
        ]

    # 2. Line Chart: Date series or ordered record progression
    line_chart_data = []
    if date_cols and primary_num:
        date_col = date_cols[0]
        date_agg: Dict[str, float] = {}
        for r in records:
            d_val = r.get(date_col)
            if not d_val or str(d_val).strip() in ["", "None", "nan", "null", "NaT"]:
                continue
            d_str = str(d_val).strip()
            num_val = r.get(primary_num)
            if num_val is not None and isinstance(num_val, (int, float)):
                if not (math.isnan(num_val) or math.isinf(num_val)):
                    date_agg[d_str] = date_agg.get(d_str, 0.0) + float(num_val)

        sorted_dates = sorted(date_agg.items(), key=lambda x: x[0])
        line_chart_data = [
            {"date": d, "value": round(v, 2), "metric": primary_num.replace("_", " ").title()}
            for d, v in sorted_dates
            if v is not None and not math.isnan(v)
        ]
    elif primary_num:
        line_chart_data = []
        idx = 1
        for r in records[:15]:
            num_val = r.get(primary_num)
            if num_val is not None and isinstance(num_val, (int, float)):
                if not (math.isnan(num_val) or math.isinf(num_val)):
                    line_chart_data.append({"index": idx, "value": round(float(num_val), 2)})
                    idx += 1

    # 3. Pie Chart: Distribution of a categorical field
    pie_chart_data = []
    pie_target_col = None
    for kw in ["status", "payment", "category", "placement", "tier"]:
        for col in cat_cols:
            if kw in col.lower():
                pie_target_col = col
                break
        if pie_target_col:
            break
    if not pie_target_col and cat_cols:
        pie_target_col = cat_cols[0]

    if pie_target_col:
        dist = kpis.get("categorical_metrics", {}).get(pie_target_col, {}).get("distribution", {})
        pie_chart_data = [
            {"name": str(k), "value": v}
            for k, v in dist.items()
            if v is not None and isinstance(v, (int, float)) and v > 0 and not math.isnan(v)
        ]

    # 4. Stat Cards
    data_quality = kpis.get("dataset_summary", {}).get("data_quality_score")
    data_quality_str = f"{data_quality}%" if data_quality is not None else "N/A"

    stat_cards = [
        {
            "title": "Total Records",
            "value": f"{cleaned_data.get('row_count', 0):,}",
            "subtitle": f"{cleaned_data.get('column_count', 0)} Columns",
            "icon": "database",
        },
        {
            "title": "Data Quality",
            "value": data_quality_str,
            "subtitle": "Complete records",
            "icon": "check-circle",
        },
    ]

    for bm in kpis.get("top_business_metrics", [])[:2]:
        stat_cards.append({
            "title": f"Total {bm.get('metric_name', 'Metric')}",
            "value": bm.get("total_display") or "N/A",
            "subtitle": f"Average: {bm.get('avg_display') or 'N/A'}",
            "icon": "trending-up",
        })

    return clean_nan_and_inf({
        "primary_metric": primary_num,
        "primary_category": primary_cat,
        "bar_chart": bar_chart_data,
        "line_chart": line_chart_data,
        "pie_chart": pie_chart_data,
        "stat_cards": stat_cards,
    })


def generate_executive_summary(
    filename: str,
    cleaned_data: Dict[str, Any],
    kpis: Dict[str, Any],
    focus_metric: Optional[str] = "",
) -> str:
    """
    Synthesizes an insightful Executive Summary in Markdown with actionable recommendations.
    """
    now_ist = datetime.now(IST).strftime("%A, %d %B %Y, %I:%M %p IST")
    rows = cleaned_data["row_count"]
    cols = cleaned_data["column_count"]
    score = kpis["dataset_summary"]["data_quality_score"]

    top_biz = kpis.get("top_business_metrics", [])
    cat_metrics = kpis.get("categorical_metrics", {})
    num_metrics = kpis.get("numeric_metrics", {})
    temporal_metrics = kpis.get("temporal_metrics", {})

    lines = [
        f"# Executive Data Analysis Report",
        f"**Source File:** `{filename}`  ",
        f"**Generated On:** {now_ist}  ",
        f"**Dataset Scope:** {rows} records | {cols} attributes | {score}% completeness score  ",
        "",
        "---",
        "",
        "## 1. Executive Overview",
    ]

    # Dynamic summary paragraph
    summary_p = f"This report delivers an automated diagnostic analysis of `{filename}`, comprising {rows} structured observations across {cols} dimensions."
    if top_biz:
        primary_kpi = top_biz[0]
        summary_p += f" Primary metric tracking shows an aggregate {primary_kpi['metric_name']} of **{primary_kpi['total_display']}**, averaging **{primary_kpi['avg_display']}** across recorded events."
    if temporal_metrics:
        first_t = list(temporal_metrics.values())[0]
        summary_p += f" The data spans from **{first_t['earliest_date']}** to **{first_t['latest_date']}**."
    lines.append(summary_p)
    lines.append("")

    # KPI Table
    lines.append("## 2. Key Performance Indicators (KPIs)")
    lines.append("| Metric / Dimension | Total Aggregate | Average / Mean | Min - Max Spread |")
    lines.append("| :--- | :--- | :--- | :--- |")

    for bm in top_biz:
        spread = f"{bm['min']} – {bm['max']}" if bm.get('min') is not None and bm.get('max') is not None else "N/A"
        tot_disp = bm.get("total_display") or "N/A"
        avg_disp = bm.get("avg_display") or "N/A"
        lines.append(f"| **{bm['metric_name']}** | `{tot_disp}` | `{avg_disp}` | {spread} |")

    for col, s in list(num_metrics.items())[:3]:
        if not any(bm.get("metric_name", "").lower() == col.replace("_", " ").lower() for bm in top_biz):
            name = col.replace("_", " ").title()
            tot_str = f"{s['total']:,}" if s.get('total') is not None else "N/A"
            avg_str = f"{s['average']:,}" if s.get('average') is not None else "N/A"
            spread = f"{s['min']} – {s['max']}" if s.get('min') is not None and s.get('max') is not None else "N/A"
            lines.append(f"| **{name}** | `{tot_str}` | `{avg_str}` | {spread} |")

    lines.append("")

    # Category Breakdowns
    if cat_metrics:
        lines.append("## 3. Categorical Distribution & Segment Leaders")
        for cat_col, cat_data in list(cat_metrics.items())[:3]:
            cat_title = cat_col.replace("_", " ").title()
            lines.append(f"### • {cat_title}")
            lines.append(f"- **Top Contributing Segment:** `{cat_data['top_value']}` ({cat_data['top_value_frequency']} occurrences, {cat_data['top_percentage']}% share)")
            lines.append(f"- **Total Distinct Segments:** {cat_data['unique_values_count']}")
            # Mini distribution
            dist_parts = [f"`{k}`: {v}" for k, v in list(cat_data["distribution"].items())[:5]]
            lines.append(f"- **Breakdown:** " + " | ".join(dist_parts))
            lines.append("")

    # Actionable Recommendations
    lines.append("## 4. Strategic Insights & Actionable Recommendations")
    lines.append("Based on the quantitative pattern analysis, the following actions are recommended:")
    lines.append("")

    # Contextual recommendation generation
    rec_count = 1
    if top_biz:
        lead_metric = top_biz[0]
        lines.append(f"{rec_count}. **Optimize High-Yield Drivers:** Prioritize operations contributing to top-tier {lead_metric['metric_name']}. Reallocate resources away from underperforming quartiles toward the peak {lead_metric['avg_display']} segments.")
        rec_count += 1

    if cat_metrics:
        lead_cat = list(cat_metrics.values())[0]
        lines.append(f"{rec_count}. **Diversify Market/Category Concentration:** `{lead_cat['top_value']}` currently accounts for {lead_cat['top_percentage']}% of total volume. Broaden outreach across secondary segments to mitigate dependency risk.")
        rec_count += 1

    lines.append(f"{rec_count}. **Data Quality & Automated Auditing:** With an overall data completeness score of {score}%, schedule recurring automated checks via AgentFlow's Scheduler to prevent missing values and detect outliers in real-time.")
    lines.append("")
    lines.append("---")
    lines.append("*Report generated autonomously by AgentFlow Report Engine (Stage 4). 100% free & local computation.*")

    return "\n".join(lines)


def export_report_files(
    base_name: str,
    cleaned_data: Dict[str, Any],
    kpis: Dict[str, Any],
    summary_md: str,
) -> Dict[str, str]:
    """
    Exports Cleaned CSV, Multi-Sheet Formatted Excel (.xlsx), and Markdown Summary (.md).
    Returns mapping of file formats to their filenames inside workspace/reports/.
    Applies strict NaN/Inf sanitization to ensure exported files are 100% valid.
    """
    clean_stem = re.sub(r"[^\w\-]", "_", Path(base_name).stem)

    csv_filename = f"{clean_stem}_cleaned.csv"
    xlsx_filename = f"{clean_stem}_analysis.xlsx"
    md_filename = f"{clean_stem}_executive_summary.md"

    csv_path = REPORTS_DIR / csv_filename
    xlsx_path = REPORTS_DIR / xlsx_filename
    md_path = REPORTS_DIR / md_filename

    # 1. Export Cleaned CSV
    records = cleaned_data.get("records", [])
    columns = cleaned_data.get("columns", [])
    if records:
        with open(csv_path, mode="w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=columns)
            writer.writeheader()
            for r in records:
                clean_row = {}
                for col in columns:
                    v = r.get(col)
                    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
                        clean_row[col] = ""
                    else:
                        clean_row[col] = v
                writer.writerow(clean_row)

    # 2. Export Markdown Summary
    with open(md_path, mode="w", encoding="utf-8") as f:
        f.write(summary_md)

    # 3. Export Beautiful Multi-Sheet Excel (.xlsx) using openpyxl
    wb = openpyxl.Workbook()

    # Sheet 1: Executive Summary Text
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    # Styling constants
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    body_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    ws_summary.cell(row=1, column=1, value="AGENTFLOW EXECUTIVE REPORT").font = title_font
    ws_summary.cell(row=2, column=1, value=f"Generated on {datetime.now(IST).strftime('%Y-%m-%d %H:%M:%S IST')} | File: {cleaned_data.get('file_name', '')}").font = subtitle_font

    ws_summary.cell(row=4, column=1, value="Dataset Overview").font = Font(name="Segoe UI", size=12, bold=True)
    ws_summary.cell(row=5, column=1, value="Total Rows").font = bold_font
    ws_summary.cell(row=5, column=2, value=cleaned_data.get("row_count", 0)).font = body_font
    ws_summary.cell(row=6, column=1, value="Total Columns").font = bold_font
    ws_summary.cell(row=6, column=2, value=cleaned_data.get("column_count", 0)).font = body_font
    ws_summary.cell(row=7, column=1, value="Completeness Score").font = bold_font
    quality_disp = kpis.get("dataset_summary", {}).get("data_quality_score")
    ws_summary.cell(row=7, column=2, value=f"{quality_disp}%" if quality_disp is not None else "N/A").font = body_font

    curr_row = 9
    ws_summary.cell(row=curr_row, column=1, value="Key Performance Indicators").font = Font(name="Segoe UI", size=12, bold=True)
    curr_row += 1
    ws_summary.cell(row=curr_row, column=1, value="Metric").font = header_font
    ws_summary.cell(row=curr_row, column=1).fill = header_fill
    ws_summary.cell(row=curr_row, column=2, value="Aggregate Total").font = header_font
    ws_summary.cell(row=curr_row, column=2).fill = header_fill
    ws_summary.cell(row=curr_row, column=3, value="Average").font = header_font
    ws_summary.cell(row=curr_row, column=3).fill = header_fill

    for bm in kpis.get("top_business_metrics", []):
        curr_row += 1
        ws_summary.cell(row=curr_row, column=1, value=bm.get("metric_name") or "Metric").font = bold_font
        ws_summary.cell(row=curr_row, column=2, value=bm.get("total_display") or "N/A").font = body_font
        ws_summary.cell(row=curr_row, column=3, value=bm.get("avg_display") or "N/A").font = body_font

    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 24
    ws_summary.column_dimensions["C"].width = 20

    # Sheet 2: Cleaned Data
    ws_data = wb.create_sheet(title="Cleaned Data")
    ws_data.views.sheetView[0].showGridLines = True

    # Write Headers
    for c_idx, col_name in enumerate(columns, start=1):
        cell = ws_data.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write Rows with safe values (None for NaN/empty)
    for r_idx, row_dict in enumerate(records, start=2):
        for c_idx, col_name in enumerate(columns, start=1):
            val = row_dict.get(col_name)
            if val is None or pd.isna(val) or (isinstance(val, float) and (math.isnan(val) or math.isinf(val))):
                cell_val = None
            elif isinstance(val, (int, float, str, bool)):
                cell_val = val
            else:
                cell_val = str(val)

            cell = ws_data.cell(row=r_idx, column=c_idx, value=cell_val)
            cell.font = body_font
            cell.border = thin_border
            if isinstance(cell_val, (int, float)):
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col[:50])
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Sheet 3: Categorical Breakdown
    if kpis.get("categorical_metrics"):
        ws_cat = wb.create_sheet(title="Categorical Breakdown")
        ws_cat.views.sheetView[0].showGridLines = True
        c_row = 1
        row_count = cleaned_data.get("row_count", 0)
        for cat_col, cat_data in kpis["categorical_metrics"].items():
            ws_cat.cell(row=c_row, column=1, value=f"{cat_col.replace('_', ' ').title()} Distribution").font = Font(name="Segoe UI", size=12, bold=True)
            c_row += 1
            ws_cat.cell(row=c_row, column=1, value="Segment / Value").font = header_font
            ws_cat.cell(row=c_row, column=1).fill = header_fill
            ws_cat.cell(row=c_row, column=2, value="Frequency").font = header_font
            ws_cat.cell(row=c_row, column=2).fill = header_fill
            ws_cat.cell(row=c_row, column=3, value="% Share").font = header_font
            ws_cat.cell(row=c_row, column=3).fill = header_fill

            for seg, count in cat_data.get("distribution", {}).items():
                c_row += 1
                pct = round((count / row_count) * 100, 1) if row_count > 0 else 0.0
                ws_cat.cell(row=c_row, column=1, value=str(seg)).font = body_font
                ws_cat.cell(row=c_row, column=2, value=count).font = body_font
                ws_cat.cell(row=c_row, column=3, value=f"{pct}%").font = body_font

            c_row += 2

        ws_cat.column_dimensions["A"].width = 28
        ws_cat.column_dimensions["B"].width = 16
        ws_cat.column_dimensions["C"].width = 16

    wb.save(xlsx_path)

    return {
        "csv": csv_filename,
        "xlsx": xlsx_filename,
        "md": md_filename,
        "csv_path": str(csv_path),
        "xlsx_path": str(xlsx_path),
        "md_path": str(md_path),
    }


def get_available_sample_datasets() -> List[Dict[str, Any]]:
    """Lists preloaded real-world sample datasets for testing."""
    samples = [
        {
            "id": "ecommerce_sales_india",
            "filename": "ecommerce_sales_india.csv",
            "title": "Indian E-Commerce Sales & Orders",
            "category": "Small Business / Retail",
            "target_audience": "SMB Owners & D2C Sellers",
            "description": "20 order records across Bengaluru, Delhi, Mumbai with category, units, UPI/COD payments, and revenue in INR.",
            "file_path": str(SAMPLES_DIR / "ecommerce_sales_india.csv"),
        },
        {
            "id": "student_placement_data",
            "filename": "student_placement_data.csv",
            "title": "College Campus Placement Analytics",
            "category": "Education / Career",
            "target_audience": "Students & Placement Cells",
            "description": "Student academic performance, CGPA, branch, tech stack, and CTC offers (in LPA) across Indian engineering colleges.",
            "file_path": str(SAMPLES_DIR / "student_placement_data.csv"),
        },
        {
            "id": "freelance_agency_expenses",
            "filename": "freelance_agency_expenses.csv",
            "title": "Freelance Agency Expense Ledger",
            "category": "Finance / Freelancing",
            "target_audience": "Freelancers & Agency Founders",
            "description": "Operational expense breakdown including software SaaS, cloud hosting, contractor payouts, and GST claims.",
            "file_path": str(SAMPLES_DIR / "freelance_agency_expenses.csv"),
        },
    ]

    # Verify existence
    result = []
    for s in samples:
        if Path(s["file_path"]).exists():
            result.append(s)
    return result
